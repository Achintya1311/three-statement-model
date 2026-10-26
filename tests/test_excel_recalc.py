"""P77: recalculate the exported workbook and compare it with Python.

``tests/test_excel_export.py`` checks that every Forecast cell is a formula;
this file checks what those formulas *evaluate to*. The workbook is
recalculated by LibreOffice headless (``soffice --convert-to xlsx`` writes
cached values for every formula), read back with ``openpyxl`` in
``data_only`` mode, and the key Forecast rows are compared with
``model.linkage.forecast_full_linkage(..., circular=False)`` -- the only mode
the workbook represents -- to within 0.01 (Rs Lakh). Skipped only when
``soffice`` is not on PATH outside CI; under CI (``CI`` set, as on GitHub
Actions, which installs ``libreoffice-calc``) a missing ``soffice`` fails
instead, so a green CI run proves the recalculation actually ran.
"""

import os
import shutil
import subprocess
from pathlib import Path

import pytest
from openpyxl import load_workbook

from model.excel_export import export_workbook
from model.historical import load_historical_statements
from model.linkage import forecast_full_linkage
from model.revenue_driver import load_segment_revenue

ROOT = Path(__file__).resolve().parent.parent
GULFOILLUB_CSV = ROOT / "data" / "GULFOILLUB" / "financials.csv"
SEGMENT_CSV = ROOT / "data" / "GULFOILLUB" / "segment_revenue.csv"
SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")
TOLERANCE = 0.01

# Forecast-sheet label -> ForecastLinkageYear field.
COMPARED_ROWS = {
    "Total revenue": "revenue",
    "EBIT": "ebit",
    "Interest expense": "interest_expense",
    "Net income": "net_income",
    "Cash (ending balance)": "cash",
    "Total debt (ending balance)": "ending_total_debt",
}

pytestmark = pytest.mark.skipif(
    SOFFICE is None and not os.environ.get("CI"), reason="LibreOffice (soffice) not installed"
)


def test_soffice_is_available_under_ci():
    if os.environ.get("CI"):
        assert SOFFICE is not None, "CI must install libreoffice-calc (P77)"


def _recalculated_forecast(tmp_path: Path, forecast_years: int, **overrides) -> dict[str, list]:
    statements = load_historical_statements(GULFOILLUB_CSV)
    segments = load_segment_revenue(GULFOILLUB_CSV, SEGMENT_CSV)
    source = export_workbook(
        "GULFOILLUB", statements, segments, forecast_years, tmp_path / "model.xlsx", **overrides
    )
    out_dir = tmp_path / "recalc"
    subprocess.run(
        [
            SOFFICE,
            f"-env:UserInstallation=file://{tmp_path / 'lo-profile'}",
            "--headless",
            "--calc",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(out_dir),
            str(source),
        ],
        check=True,
        capture_output=True,
        timeout=180,
    )
    ws = load_workbook(out_dir / "model.xlsx", data_only=True)["Forecast"]
    rows = {}
    for row in ws.iter_rows(min_row=1):
        label = row[0].value
        if isinstance(label, str):
            # Column B is Base; forecast years start at column C.
            rows[label] = [cell.value for cell in row[2 : 2 + forecast_years]]
    return rows


@pytest.mark.parametrize(
    "overrides",
    [
        {},
        {"tax_rate": 0.30, "interest_rate": 0.09, "dividend_payout": 0.25, "ebitda_margin": 0.12},
    ],
    ids=["defaults", "overrides"],
)
def test_recalculated_workbook_matches_non_circular_linkage(tmp_path, overrides):
    forecast_years = 3
    rows = _recalculated_forecast(tmp_path, forecast_years, **overrides)
    expected = forecast_full_linkage(
        load_historical_statements(GULFOILLUB_CSV),
        load_segment_revenue(GULFOILLUB_CSV, SEGMENT_CSV),
        forecast_years,
        circular=False,
        **overrides,
    )
    assert rows["Fiscal year"] == [year.fiscal_year for year in expected]
    for label, field in COMPARED_ROWS.items():
        for got, year in zip(rows[label], expected):
            want = getattr(year, field)
            assert isinstance(got, (int, float)), f"{label} {year.fiscal_year}: no value ({got!r})"
            assert abs(got - want) <= TOLERANCE, f"{label} {year.fiscal_year}: workbook {got} vs python {want}"
    assert rows["Balance check"] == ["PASS" if year.balance_check_passed else "FAIL" for year in expected]
